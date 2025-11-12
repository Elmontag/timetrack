from __future__ import annotations

import datetime as dt
import io
import re
import shutil
import uuid
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

from zoneinfo import ZoneInfo

from fastapi import HTTPException, status
from openpyxl import Workbook
from PyPDF2 import PdfMerger
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from sqlalchemy import and_, func, or_
from sqlalchemy.orm import Session

from .config import settings
from .models import (
    CalendarEvent,
    DaySummary,
    ExportRecord,
    Holiday,
    LeaveEntry,
    TravelDocument,
    TravelTrip,
    WorkSession,
    WorkSubtrack,
)
from .state import RuntimeState
from .utils import normalize_calendar_identifier, normalize_calendar_selection

try:  # pragma: no cover - optional dependency
    from caldav import DAVClient
    from caldav.elements import dav
    from caldav.lib import error as caldav_error
except ImportError as exc:  # pragma: no cover - caldav optional
    DAVClient = None  # type: ignore[assignment]
    dav = None  # type: ignore[assignment]
    caldav_error = None  # type: ignore[assignment]
    _CALDAV_IMPORT_ERROR = exc
else:  # pragma: no cover - only set when caldav is available
    _CALDAV_IMPORT_ERROR = None


UTC = dt.timezone.utc
LOCAL_TZ = ZoneInfo(settings.timezone)


AUTO_SESSION_COMMENT_PREFIX = "Automatisch aus Termin: "

CALENDAR_EVENT_STATUSES: Set[str] = {"pending", "attended", "absent", "cancelled"}

TRAVEL_WORKFLOW_STATES: Dict[str, str] = {
    "request_draft": "Dienstreiseantrag",
    "requested": "Dienstreise beantragt",
    "settlement": "Reise wird abgerechnet",
    "settled": "Reise abgerechnet",
    "reimbursed": "Kostenerstattung erhalten",
}

TRAVEL_DOCUMENT_TYPES: Set[str] = {
    "Rechnung",
    "Antrag",
    "Beleg",
    "Reisekostenabrechnung",
    "Sonstige Unterlagen",
    "Anschreiben",
}

TRAVEL_DATASET_DOCUMENT_TYPES: Set[str] = {"Rechnung", "Beleg", "Reisekostenabrechnung", "Anschreiben"}

SIGNABLE_DOCUMENT_TYPES: Set[str] = {"Antrag", "Reisekostenabrechnung"}

UNSET: Any = object()


def _normalize_collection_label(value: Optional[str]) -> Optional[str]:
    if value is None:
        return None
    normalized = value.strip()
    return normalized or None


def _now() -> dt.datetime:
    return dt.datetime.now(UTC)


def _ensure_utc(value: dt.datetime) -> dt.datetime:
    if value.tzinfo is None:
        value = value.replace(tzinfo=LOCAL_TZ)
    return value.astimezone(UTC)


def _day_bounds(day: dt.date) -> Tuple[dt.datetime, dt.datetime]:
    start_local = dt.datetime.combine(day, dt.time.min, tzinfo=LOCAL_TZ)
    end_local = start_local + dt.timedelta(days=1)
    return start_local.astimezone(UTC), end_local.astimezone(UTC)


def _build_caldav_client(state: RuntimeState, *, strict: bool = False):
    if DAVClient is None:
        if strict and _CALDAV_IMPORT_ERROR is not None:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail=f"CalDAV-Bibliothek nicht verfügbar: {_CALDAV_IMPORT_ERROR}",
            )
        if strict:
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="CalDAV support is not installed",
            )
        return None
    if not (state.caldav_url and state.caldav_user and state.caldav_password):
        if strict:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="CalDAV-Zugangsdaten sind unvollständig",
            )
        return None
    try:
        return DAVClient(
            url=state.caldav_url,
            username=state.caldav_user,
            password=state.caldav_password,
        )
    except Exception as exc:  # pragma: no cover - network errors
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail="CalDAV-Client konnte nicht initialisiert werden",
        ) from exc


def _calendar_matches_selection(selection: Iterable[str], *candidates: Optional[Any]) -> bool:
    normalized = {value.casefold() for value in normalize_calendar_selection(selection)}
    if not normalized:
        return False
    for candidate in candidates:
        candidate_id = normalize_calendar_identifier(candidate)
        if candidate_id and candidate_id.casefold() in normalized:
            return True
    return False


def _fetch_caldav_occurrences(
    calendar: Any,
    range_start: dt.datetime,
    range_end: dt.datetime,
) -> Iterable[Any]:
    def _strip_timezone(value: dt.datetime) -> dt.datetime:
        if value.tzinfo is None:
            return value
        return value.replace(tzinfo=None)

    attempts: List[Tuple[dt.datetime, dt.datetime, Optional[bool]]] = [
        (range_start, range_end, True),
        (range_start, range_end, False),
    ]

    naive_start = _strip_timezone(range_start)
    naive_end = _strip_timezone(range_end)
    if naive_start is not range_start or naive_end is not range_end:
        attempts.extend(
            [
                (naive_start, naive_end, True),
                (naive_start, naive_end, False),
            ]
        )

    last_error: Optional[Exception] = None
    date_search = getattr(calendar, "date_search", None)
    if callable(date_search):
        for start, end, expand in attempts:
            try:
                if expand is None:
                    occurrences = date_search(start, end)
                else:
                    occurrences = date_search(start, end, expand=expand)
            except Exception as exc:  # pragma: no cover - dependent on remote server
                last_error = exc
                continue
            if occurrences is not None:
                return occurrences

    events_method = getattr(calendar, "events", None)
    if callable(events_method):
        try:
            occurrences = events_method()
        except Exception as exc:  # pragma: no cover - dependent on remote server
            last_error = exc
        else:
            if occurrences is not None:
                return occurrences

    if last_error is not None:
        raise last_error
    return []


def _coerce_ical_datetime(value: Any) -> Optional[dt.datetime]:
    if value is None:
        return None
    attr = getattr(value, "dt", value)
    if isinstance(attr, dt.datetime):
        if attr.tzinfo is None:
            return attr.replace(tzinfo=LOCAL_TZ)
        return attr
    if isinstance(attr, dt.date):
        return dt.datetime.combine(attr, dt.time.min, tzinfo=LOCAL_TZ)
    return None


def fetch_caldav_calendars(state: RuntimeState) -> List[Dict[str, str]]:
    client = _build_caldav_client(state, strict=True)
    try:
        principal = client.principal()
        calendars = principal.calendars()
    except Exception as exc:  # pragma: no cover - remote errors
        auth_error = getattr(caldav_error, "AuthorizationError", None)
        if auth_error and isinstance(exc, auth_error):  # type: ignore[arg-type]
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="CalDAV-Anmeldung fehlgeschlagen") from exc
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail="CalDAV-Kalender konnten nicht geladen werden") from exc

    results: List[Dict[str, str]] = []
    for calendar in calendars:
        calendar_id_raw = getattr(calendar, "url", None)
        calendar_id = normalize_calendar_identifier(calendar_id_raw)
        display_name: Optional[str] = None
        if dav is not None and hasattr(calendar, "get_properties"):
            try:
                props = calendar.get_properties([dav.DisplayName()])  # type: ignore[attr-defined]
                display_value = props.get(dav.DisplayName()) if props else None  # type: ignore[attr-defined]
                if display_value:
                    display_name = str(display_value)
            except Exception:  # pragma: no cover - best effort only
                display_name = None
        if not display_name:
            name_attr = getattr(calendar, "name", None)
            if name_attr is not None:
                display_name = str(name_attr)
        if not display_name and calendar_id:
            display_name = calendar_id.split("/")[-1]
        if not calendar_id and display_name:
            calendar_id = normalize_calendar_identifier(display_name)
        if not calendar_id:
            continue
        results.append({"id": calendar_id, "name": display_name or calendar_id})
    return results


def sync_caldav_events(
    db: Session,
    state: RuntimeState,
    start_date: Optional[dt.date],
    end_date: Optional[dt.date],
) -> None:
    client = _build_caldav_client(state)
    selected = state.caldav_selected_calendars
    normalized_selected = {value for value in normalize_calendar_selection(selected)}
    if client is None or not normalized_selected:
        return

    start_day = start_date or dt.date.today()
    end_day = end_date or start_day
    range_start, _ = _day_bounds(start_day)
    _, range_end = _day_bounds(end_day)

    try:
        principal = client.principal()
        calendars = principal.calendars()
    except Exception as exc:  # pragma: no cover - remote errors
        auth_error = getattr(caldav_error, "AuthorizationError", None)
        if auth_error and isinstance(exc, auth_error):  # type: ignore[arg-type]
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="CalDAV-Anmeldung fehlgeschlagen") from exc
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail="CalDAV-Synchronisation fehlgeschlagen") from exc

    existing_by_uid: Dict[Tuple[Optional[str], Optional[str], Optional[str]], CalendarEvent] = {}
    existing_without_uid: Dict[Tuple[Optional[str], dt.datetime, dt.datetime, str], CalendarEvent] = {}
    duplicates: List[CalendarEvent] = []

    existing_query = db.query(CalendarEvent).filter(
        or_(
            CalendarEvent.calendar_identifier.is_(None),
            CalendarEvent.calendar_identifier.in_(normalized_selected),
        )
    )

    for stored in existing_query.all():
        if stored.external_id:
            key = (
                stored.calendar_identifier,
                stored.external_id,
                stored.recurrence_id or "",
            )
            if key in existing_by_uid:
                duplicates.append(stored)
                continue
            existing_by_uid[key] = stored
        else:
            key = (
                stored.calendar_identifier,
                _ensure_utc(stored.start_time),
                _ensure_utc(stored.end_time),
                stored.title,
            )
            if key in existing_without_uid:
                duplicates.append(stored)
                continue
            existing_without_uid[key] = stored

    updated = False
    for duplicate in duplicates:
        db.delete(duplicate)
        updated = True

    for calendar in calendars:
        calendar_id_raw = getattr(calendar, "url", None)
        calendar_id = normalize_calendar_identifier(calendar_id_raw)
        calendar_name_raw = getattr(calendar, "name", None)
        calendar_name = normalize_calendar_identifier(calendar_name_raw)
        if not _calendar_matches_selection(selected, calendar_id, calendar_name):
            continue
        if calendar_id and calendar_id in normalized_selected:
            event_calendar_identifier = calendar_id
        elif calendar_name and calendar_name in normalized_selected:
            event_calendar_identifier = calendar_name
        else:
            event_calendar_identifier = calendar_id or calendar_name or "caldav"
        try:
            occurrences = _fetch_caldav_occurrences(calendar, range_start, range_end)
        except Exception as exc:  # pragma: no cover - propagate as HTTP error
            identifier = calendar_name or calendar_id or "unbekannt"
            raise HTTPException(
                status_code=status.HTTP_502_BAD_GATEWAY,
                detail=f"CalDAV-Kalender '{identifier}' konnte nicht synchronisiert werden",
            ) from exc
        for occurrence in occurrences:
            component = getattr(occurrence, "icalendar_component", None)
            vevent = _extract_vevent(component)
            if vevent is None:
                continue
            summary = _ical_to_string(vevent.get("summary")) or "Unbenannter Termin"
            dtstart = _coerce_ical_datetime(vevent.get("dtstart"))
            dtend = _coerce_ical_datetime(vevent.get("dtend"))
            if dtstart is None:
                continue
            start_utc = _ensure_utc(dtstart)
            if dtend is None:
                dtend = dtstart + dt.timedelta(hours=1)
            end_utc = _ensure_utc(dtend)
            if start_utc < range_start or start_utc > range_end:
                continue
            location = _ical_to_string(vevent.get("location"))
            description = _ical_to_string(vevent.get("description"))
            external_id = _ical_to_string(vevent.get("uid"))
            recurrence_id = _ical_to_string(vevent.get("recurrence_id"))
            attendees = _extract_attendees(vevent)

            existing_event: Optional[CalendarEvent] = None
            if external_id:
                existing_event = existing_by_uid.get(
                    (event_calendar_identifier, external_id, recurrence_id or "")
                )

            fallback_key = (
                event_calendar_identifier,
                start_utc,
                end_utc,
                summary,
            )
            if existing_event is None:
                existing_event = existing_without_uid.get(fallback_key)

            if existing_event is not None:
                old_fallback_key: Optional[
                    Tuple[Optional[str], dt.datetime, dt.datetime, str]
                ] = None
                if existing_event.external_id is None:
                    old_fallback_key = (
                        existing_event.calendar_identifier,
                        _ensure_utc(existing_event.start_time),
                        _ensure_utc(existing_event.end_time),
                        existing_event.title,
                    )

                changed = False
                if existing_event.title != summary:
                    existing_event.title = summary
                    changed = True
                if _ensure_utc(existing_event.start_time) != start_utc:
                    existing_event.start_time = start_utc
                    changed = True
                if _ensure_utc(existing_event.end_time) != end_utc:
                    existing_event.end_time = end_utc
                    changed = True
                if existing_event.location != location:
                    existing_event.location = location
                    changed = True
                if existing_event.description != description:
                    existing_event.description = description
                    changed = True
                if existing_event.calendar_identifier != event_calendar_identifier:
                    existing_event.calendar_identifier = event_calendar_identifier
                    changed = True
                if existing_event.external_id != external_id:
                    existing_event.external_id = external_id
                    changed = True
                if existing_event.recurrence_id != recurrence_id:
                    existing_event.recurrence_id = recurrence_id
                    changed = True
                if existing_event.attendees != attendees:
                    existing_event.attendees = attendees
                    changed = True

                if existing_event.external_id:
                    key = (
                        existing_event.calendar_identifier,
                        existing_event.external_id,
                        existing_event.recurrence_id or "",
                    )
                    existing_by_uid[key] = existing_event
                else:
                    if old_fallback_key and old_fallback_key != fallback_key:
                        existing_without_uid.pop(old_fallback_key, None)
                    existing_without_uid[fallback_key] = existing_event

                if changed:
                    updated = True
                continue

            event = CalendarEvent(
                title=summary,
                start_time=start_utc,
                end_time=end_utc,
                location=location,
                description=description,
                participated=False,
                calendar_identifier=event_calendar_identifier,
                external_id=external_id,
                recurrence_id=recurrence_id,
                attendees=attendees,
            )
            db.add(event)
            if external_id:
                key = (event_calendar_identifier, external_id, recurrence_id or "")
                existing_by_uid[key] = event
            else:
                existing_without_uid[fallback_key] = event
            updated = True

    if updated:
        db.commit()


def _extract_vevent(component: Any) -> Optional[Any]:
    """Return the VEVENT component from a CalDAV occurrence."""

    if component is None:
        return None

    candidates: List[Any] = [component]

    direct = getattr(component, "vevent", None)
    if direct is not None:
        candidates.append(direct)

    subcomponents = getattr(component, "subcomponents", None)
    if subcomponents:
        candidates.extend(subcomponents)

    for candidate in candidates:
        if candidate is None:
            continue

        name = getattr(candidate, "name", None)
        if isinstance(name, str) and name.upper() == "VEVENT":
            return candidate

        get = getattr(candidate, "get", None)
        if callable(get):
            try:
                if get("dtstart") is not None or get("summary") is not None:
                    return candidate
            except Exception:  # pragma: no cover - defensive guard
                continue

    return None


def _ical_to_string(value: Any) -> Optional[str]:
    if value is None:
        return None
    to_ical = getattr(value, "to_ical", None)
    if callable(to_ical):
        try:
            rendered = to_ical()
        except Exception:  # pragma: no cover - defensive conversion
            rendered = None
        else:
            if isinstance(rendered, bytes):
                return rendered.decode("utf-8", errors="ignore")
            return str(rendered)
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="ignore")
    return str(value)


def _extract_attendees(vevent: Any) -> List[str]:
    attendees_raw = vevent.get("attendee")
    if not attendees_raw:
        return []
    if not isinstance(attendees_raw, list):
        attendees_raw = [attendees_raw]
    attendees: List[str] = []
    for entry in attendees_raw:
        name: Optional[str] = None
        params = getattr(entry, "params", None)
        if params:
            cn_value = params.get("CN")
            if isinstance(cn_value, list):
                cn_value = cn_value[0] if cn_value else None
            if cn_value is not None:
                name = _ical_to_string(cn_value)
        if not name:
            name = _ical_to_string(entry)
        if name:
            attendees.append(name)
    return attendees


def _expected_daily_seconds(state: RuntimeState) -> int:
    if state.expected_daily_hours is not None:
        return int(state.expected_daily_hours * 3600)
    if state.expected_weekly_hours is not None:
        return int(state.expected_weekly_hours / 5 * 3600)
    return int(settings.expected_daily_hours * 3600)


def _from_db_datetime(value: Optional[dt.datetime]) -> Optional[dt.datetime]:
    if value is None:
        return None
    if value.tzinfo is None:
        return value.replace(tzinfo=UTC)
    return value.astimezone(UTC)


def get_active_session(db: Session) -> Optional[WorkSession]:
    return (
        db.query(WorkSession)
        .filter(WorkSession.status.in_(["active", "paused"]))
        .order_by(WorkSession.start_time.desc())
        .first()
    )


def start_session(
    db: Session,
    project: Optional[str],
    tags: List[str],
    comment: Optional[str],
    start_time: Optional[dt.datetime] = None,
) -> WorkSession:
    active = get_active_session(db)
    if active:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Active session already exists")
    session = WorkSession(
        start_time=_ensure_utc(start_time) if start_time else _now(),
        project=project,
        tags=tags,
        comment=comment,
    )
    db.add(session)
    db.commit()
    db.refresh(session)
    return session


def pause_or_resume_session(db: Session) -> WorkSession:
    session = get_active_session(db)
    if not session:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="No active session")
    now = _now()
    if session.status == "paused":
        session.mark_resumed(now)
        action = "resumed"
    else:
        session.mark_paused(now)
        action = "paused"
    db.add(session)
    db.commit()
    db.refresh(session)
    session._last_action = action  # type: ignore[attr-defined]
    return session


def stop_session(db: Session, state: RuntimeState, comment: Optional[str] = None) -> WorkSession:
    session = get_active_session(db)
    if not session:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="No active session")
    now = _now()
    if comment:
        session.comment = comment
    session.mark_stopped(now)
    db.add(session)
    db.commit()
    db.refresh(session)
    start_local = _from_db_datetime(session.start_time).astimezone(LOCAL_TZ)
    update_day_summary(db, start_local.date(), state)
    return session


def create_manual_session(
    db: Session,
    state: RuntimeState,
    start_time: dt.datetime,
    end_time: dt.datetime,
    project: Optional[str],
    tags: List[str],
    comment: Optional[str],
) -> WorkSession:
    if end_time <= start_time:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="End time must be after start time")
    start_time_utc = _ensure_utc(start_time)
    end_time_utc = _ensure_utc(end_time)
    duration = end_time_utc - start_time_utc
    total_seconds = int(duration.total_seconds())
    session = WorkSession(
        start_time=start_time_utc,
        stop_time=end_time_utc,
        status="stopped",
        project=project,
        tags=tags,
        comment=comment,
        paused_duration=0,
        total_seconds=total_seconds,
    )
    db.add(session)
    db.commit()
    db.refresh(session)
    update_day_summary(db, start_time_utc.astimezone(LOCAL_TZ).date(), state)
    return session


def _auto_session_comment(title: str, note: Optional[str] = None) -> str:
    base = f"{AUTO_SESSION_COMMENT_PREFIX}{title}"
    if note:
        return f"{base} – {note}"
    return base


def _ensure_session_for_span(
    db: Session,
    state: RuntimeState,
    start_time: Optional[dt.datetime],
    end_time: Optional[dt.datetime],
    title: str,
    project: Optional[str],
    tags: List[str],
    note: Optional[str],
) -> Optional[WorkSession]:
    if start_time is None or end_time is None:
        return None
    start_utc = _ensure_utc(start_time)
    end_utc = _ensure_utc(end_time)
    existing = (
        db.query(WorkSession)
        .filter(
            WorkSession.start_time == start_utc,
            WorkSession.stop_time == end_utc,
        )
        .one_or_none()
    )
    desired_comment = _auto_session_comment(title, note)
    if existing:
        if (existing.comment or "").startswith(AUTO_SESSION_COMMENT_PREFIX) and existing.comment != desired_comment:
            existing.comment = desired_comment
            db.add(existing)
        return existing
    return create_manual_session(
        db,
        state,
        start_utc,
        end_utc,
        project,
        tags or [],
        desired_comment,
    )


def _ensure_session_for_subtrack(
    db: Session,
    state: RuntimeState,
    subtrack: WorkSubtrack,
    source_title: Optional[str] = None,
) -> Optional[WorkSession]:
    title = source_title or subtrack.title
    return _ensure_session_for_span(
        db,
        state,
        _from_db_datetime(subtrack.start_time),
        _from_db_datetime(subtrack.end_time),
        title,
        subtrack.project,
        list(subtrack.tags or []),
        subtrack.note,
    )


def _remove_auto_session(
    db: Session,
    state: RuntimeState,
    start_time: Optional[dt.datetime],
    end_time: Optional[dt.datetime],
    title: str,
    note: Optional[str],
) -> Optional[dt.date]:
    start_utc = _from_db_datetime(start_time)
    end_utc = _from_db_datetime(end_time)
    if start_utc is None or end_utc is None:
        return None
    auto_comment = _auto_session_comment(title, note)
    session = (
        db.query(WorkSession)
        .filter(
            WorkSession.start_time == start_utc,
            WorkSession.stop_time == end_utc,
            WorkSession.comment == auto_comment,
        )
        .one_or_none()
    )
    if session is None:
        return None
    day = start_utc.astimezone(LOCAL_TZ).date()
    db.delete(session)
    update_day_summary(db, day, state)
    return day


def list_sessions_for_day(db: Session, day: dt.date) -> List[WorkSession]:
    start, end = _day_bounds(day)
    sessions = (
        db.query(WorkSession)
        .filter(and_(WorkSession.start_time >= start, WorkSession.start_time < end))
        .order_by(WorkSession.start_time.asc())
        .all()
    )
    return sessions


def _list_holidays(db: Session, start_day: dt.date, end_day: dt.date) -> List[Holiday]:
    return (
        db.query(Holiday)
        .filter(and_(Holiday.day >= start_day, Holiday.day <= end_day))
        .order_by(Holiday.day.asc())
        .all()
    )


def _holiday_lookup(db: Session, start_day: dt.date, end_day: dt.date) -> Dict[dt.date, Holiday]:
    holidays = _list_holidays(db, start_day, end_day)
    return {holiday.day: holiday for holiday in holidays}


def _leave_day_map(db: Session, start_day: dt.date, end_day: dt.date) -> Dict[dt.date, Set[str]]:
    leaves = (
        db.query(LeaveEntry)
        .filter(
            and_(
                LeaveEntry.end_date >= start_day,
                LeaveEntry.start_date <= end_day,
            )
        )
        .all()
    )
    mapping: Dict[dt.date, Set[str]] = defaultdict(set)
    for entry in leaves:
        current = max(entry.start_date, start_day)
        final = min(entry.end_date, end_day)
        while current <= final:
            mapping[current].add(entry.type)
            current += dt.timedelta(days=1)
    return mapping


def _calculate_daily_context(
    day: dt.date,
    state: RuntimeState,
    leave_types: Set[str],
    holiday: Optional[Holiday],
) -> Tuple[int, int, int, int, bool, bool]:
    base_expected = _expected_daily_seconds(state)
    is_weekend = day.weekday() >= 5
    is_holiday = holiday is not None
    expected = base_expected
    if is_weekend or is_holiday or "sick" in leave_types:
        expected = 0
    vacation_seconds = 0
    if "vacation" in leave_types and not (is_weekend or is_holiday):
        vacation_seconds = base_expected
    sick_seconds = base_expected if "sick" in leave_types else 0
    return expected, base_expected, vacation_seconds, sick_seconds, is_weekend, is_holiday


def _aggregate_sessions(
    db: Session,
    start_day: dt.date,
    end_day: dt.date,
) -> Dict[dt.date, Dict[str, int]]:
    range_start, _ = _day_bounds(start_day)
    _, range_end = _day_bounds(end_day)
    sessions = (
        db.query(WorkSession)
        .filter(
            and_(
                WorkSession.start_time >= range_start,
                WorkSession.start_time < range_end,
                WorkSession.status == "stopped",
            )
        )
        .all()
    )
    totals: Dict[dt.date, Dict[str, int]] = defaultdict(lambda: {"work": 0, "pause": 0})
    for session in sessions:
        local_start = _from_db_datetime(session.start_time).astimezone(LOCAL_TZ)
        day = local_start.date()
        if day < start_day or day > end_day:
            continue
        total_seconds = session.total_seconds
        if total_seconds is None and session.stop_time:
            start_utc = _from_db_datetime(session.start_time)
            stop_utc = _from_db_datetime(session.stop_time)
            total_seconds = int((stop_utc - start_utc).total_seconds()) - (session.paused_duration or 0)
        totals[day]["work"] += max(total_seconds or 0, 0)
        totals[day]["pause"] += session.paused_duration or 0
    return totals


def update_day_summary(db: Session, day: dt.date, state: RuntimeState) -> DaySummary:
    totals = _aggregate_sessions(db, day, day)
    leave_map = _leave_day_map(db, day, day)
    holiday_map = _holiday_lookup(db, day, day)
    leave_types = leave_map.get(day, set())
    expected, _base_expected, vacation_seconds, _sick_seconds, _is_weekend, _is_holiday = _calculate_daily_context(
        day, state, leave_types, holiday_map.get(day)
    )
    summary = db.query(DaySummary).filter(DaySummary.day == day).one_or_none() or DaySummary(day=day)
    day_totals = totals.get(day, {"work": 0, "pause": 0})
    summary.work_seconds = day_totals["work"]
    summary.pause_seconds = day_totals["pause"]
    summary.overtime_seconds = summary.work_seconds + vacation_seconds - expected
    summary.updated_at = _now()
    db.add(summary)
    db.commit()
    db.refresh(summary)
    return summary


def range_day_summaries(db: Session, start_day: dt.date, end_day: dt.date, state: RuntimeState) -> List[Dict[str, Any]]:
    totals = _aggregate_sessions(db, start_day, end_day)
    leave_map = _leave_day_map(db, start_day, end_day)
    holiday_map = _holiday_lookup(db, start_day, end_day)
    summaries: List[Dict[str, Any]] = []
    current = start_day
    while current <= end_day:
        totals_for_day = totals.get(current, {"work": 0, "pause": 0})
        leave_types = leave_map.get(current, set())
        expected, base_expected, vacation_seconds, sick_seconds, is_weekend, is_holiday = _calculate_daily_context(
            current, state, leave_types, holiday_map.get(current)
        )
        work_seconds = totals_for_day["work"]
        pause_seconds = totals_for_day["pause"]
        overtime_seconds = work_seconds + vacation_seconds - expected
        summaries.append(
            {
                "day": current,
                "work_seconds": work_seconds,
                "pause_seconds": pause_seconds,
                "overtime_seconds": overtime_seconds,
                "expected_seconds": expected,
                "vacation_seconds": vacation_seconds,
                "sick_seconds": sick_seconds,
                "is_weekend": is_weekend,
                "is_holiday": is_holiday,
                "holiday_name": holiday_map.get(current).name if is_holiday else None,
                "leave_types": sorted(leave_types),
                "baseline_expected_seconds": base_expected,
            }
        )
        current += dt.timedelta(days=1)
    return summaries


def update_session(
    db: Session,
    state: RuntimeState,
    session_id: int,
    changes: Dict[str, Any],
) -> WorkSession:
    session = db.get(WorkSession, session_id)
    if not session:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")
    if session.status != "stopped":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only stopped sessions can be edited")

    old_day = _from_db_datetime(session.start_time).astimezone(LOCAL_TZ).date()

    if "start_time" in changes and changes["start_time"] is not None:
        session.start_time = _ensure_utc(changes["start_time"])
    if "end_time" in changes and changes["end_time"] is not None:
        session.stop_time = _ensure_utc(changes["end_time"])
    if session.stop_time is None:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Stop time is required")
    if session.start_time >= session.stop_time:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Start must be before end")

    if "comment" in changes:
        session.comment = changes["comment"]
    if "project" in changes:
        session.project = changes["project"]
    if "tags" in changes:
        session.tags = changes["tags"] or []

    start_utc = _from_db_datetime(session.start_time)
    stop_utc = _from_db_datetime(session.stop_time)
    duration = stop_utc - start_utc
    session.total_seconds = max(int(duration.total_seconds()) - (session.paused_duration or 0), 0)
    session.last_pause_start = None
    session.status = "stopped"

    new_day = _from_db_datetime(session.start_time).astimezone(LOCAL_TZ).date()

    db.add(session)
    db.commit()
    db.refresh(session)

    affected_days = {old_day, new_day}
    for day in affected_days:
        update_day_summary(db, day, state)

    return session


def delete_session(db: Session, state: RuntimeState, session_id: int) -> None:
    session = db.get(WorkSession, session_id)
    if not session:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")
    if session.status != "stopped":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Active sessions cannot be deleted")

    day = _from_db_datetime(session.start_time).astimezone(LOCAL_TZ).date()
    db.delete(session)
    db.commit()
    update_day_summary(db, day, state)


def create_subtrack(
    db: Session,
    state: RuntimeState,
    day: dt.date,
    title: str,
    start_time: Optional[dt.datetime],
    end_time: Optional[dt.datetime],
    project: Optional[str],
    tags: List[str],
    note: Optional[str],
) -> WorkSubtrack:
    start_time_utc = _ensure_utc(start_time) if start_time else None
    end_time_utc = _ensure_utc(end_time) if end_time else None
    if start_time_utc and end_time_utc and end_time_utc <= start_time_utc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Subtrack end must be after start")
    subtrack = WorkSubtrack(
        day=day,
        title=title,
        start_time=start_time_utc,
        end_time=end_time_utc,
        project=project,
        tags=tags,
        note=note,
    )
    db.add(subtrack)
    db.commit()
    db.refresh(subtrack)
    _ensure_session_for_subtrack(db, state, subtrack)
    update_day_summary(db, day, state)
    return subtrack


def list_subtracks(db: Session, day: dt.date) -> List[WorkSubtrack]:
    return (
        db.query(WorkSubtrack)
        .filter(WorkSubtrack.day == day)
        .order_by(WorkSubtrack.start_time.asc(), WorkSubtrack.created_at.asc())
        .all()
    )


def update_subtrack(
    db: Session,
    state: RuntimeState,
    subtrack_id: int,
    changes: Dict[str, Any],
) -> WorkSubtrack:
    subtrack = db.get(WorkSubtrack, subtrack_id)
    if not subtrack:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Subtrack not found")

    original_day = subtrack.day
    _remove_auto_session(db, state, subtrack.start_time, subtrack.end_time, subtrack.title, subtrack.note)

    if "day" in changes and changes["day"] is not None:
        subtrack.day = changes["day"]
    if "title" in changes and changes["title"] is not None:
        subtrack.title = changes["title"]
    if "start_time" in changes:
        value = changes["start_time"]
        subtrack.start_time = _ensure_utc(value) if value else None
    if "end_time" in changes:
        value = changes["end_time"]
        subtrack.end_time = _ensure_utc(value) if value else None
    if "project" in changes:
        subtrack.project = changes["project"]
    if "tags" in changes and changes["tags"] is not None:
        subtrack.tags = changes["tags"]
    if "note" in changes:
        subtrack.note = changes["note"]

    db.add(subtrack)
    db.commit()
    db.refresh(subtrack)

    _ensure_session_for_subtrack(db, state, subtrack)

    affected_days: Set[dt.date] = {original_day, subtrack.day}
    for day in affected_days:
        update_day_summary(db, day, state)
    return subtrack


def delete_subtrack(db: Session, state: RuntimeState, subtrack_id: int) -> None:
    subtrack = db.get(WorkSubtrack, subtrack_id)
    if not subtrack:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Subtrack not found")

    day = subtrack.day
    _remove_auto_session(db, state, subtrack.start_time, subtrack.end_time, subtrack.title, subtrack.note)

    db.delete(subtrack)
    db.commit()

    update_day_summary(db, day, state)


def create_leave(
    db: Session,
    start_date: dt.date,
    end_date: dt.date,
    leave_type: str,
    comment: Optional[str],
    approved: bool,
) -> Dict[str, Any]:
    if end_date < start_date:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="End date must be after start date")
    entry = LeaveEntry(
        start_date=start_date,
        end_date=end_date,
        type=leave_type,
        comment=comment,
        approved=1 if approved else 0,
    )
    db.add(entry)
    db.commit()
    db.refresh(entry)
    holiday_lookup = _holiday_lookup(db, entry.start_date, entry.end_date)
    return {
        "id": entry.id,
        "start_date": entry.start_date,
        "end_date": entry.end_date,
        "type": entry.type,
        "comment": entry.comment,
        "approved": bool(entry.approved),
        "day_count": _effective_leave_days(entry, holiday_lookup),
    }


def _effective_leave_days(entry: LeaveEntry, holiday_lookup: Dict[dt.date, Holiday]) -> float:
    count = 0.0
    current = entry.start_date
    while current <= entry.end_date:
        if current.weekday() >= 5 or current in holiday_lookup:
            current += dt.timedelta(days=1)
            continue
        count += 1
        current += dt.timedelta(days=1)
    return count


def list_leaves(db: Session, start_date: Optional[dt.date], end_date: Optional[dt.date], leave_type: Optional[str]) -> List[Dict[str, Any]]:
    query = db.query(LeaveEntry)
    if start_date:
        query = query.filter(LeaveEntry.start_date >= start_date)
    if end_date:
        query = query.filter(LeaveEntry.end_date <= end_date)
    if leave_type:
        query = query.filter(LeaveEntry.type == leave_type)
    entries = query.order_by(LeaveEntry.start_date.asc()).all()
    if not entries:
        return []
    range_start = min(entry.start_date for entry in entries)
    range_end = max(entry.end_date for entry in entries)
    holiday_lookup = _holiday_lookup(db, range_start, range_end)
    results: List[Dict[str, Any]] = []
    for entry in entries:
        day_count = _effective_leave_days(entry, holiday_lookup)
        results.append(
            {
                "id": entry.id,
                "start_date": entry.start_date,
                "end_date": entry.end_date,
                "type": entry.type,
                "comment": entry.comment,
                "approved": bool(entry.approved),
                "day_count": day_count,
            }
        )
    return results


def list_holidays(db: Session, start_date: Optional[dt.date], end_date: Optional[dt.date]) -> List[Holiday]:
    query = db.query(Holiday)
    if start_date:
        query = query.filter(Holiday.day >= start_date)
    if end_date:
        query = query.filter(Holiday.day <= end_date)
    return query.order_by(Holiday.day.asc()).all()


def create_holiday(db: Session, day: dt.date, name: str, source: str = "manual") -> Holiday:
    existing = db.query(Holiday).filter(Holiday.day == day).one_or_none()
    if existing:
        existing.name = name
        existing.source = source
        db.add(existing)
        db.commit()
        db.refresh(existing)
        return existing
    holiday = Holiday(day=day, name=name, source=source)
    db.add(holiday)
    db.commit()
    db.refresh(holiday)
    return holiday


def delete_holiday(db: Session, holiday_id: int) -> None:
    holiday = db.get(Holiday, holiday_id)
    if not holiday:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Holiday not found")
    db.delete(holiday)
    db.commit()


def _parse_ics_holidays(content: str) -> List[Tuple[dt.date, str]]:
    if not content.strip():
        return []
    normalized = content.replace("\r\n", "\n")
    normalized = re.sub(r"\n[ \t]", "", normalized)
    events: Dict[dt.date, str] = {}
    inside = False
    current_date: Optional[dt.date] = None
    current_summary: Optional[str] = None
    for raw_line in normalized.splitlines():
        line = raw_line.strip()
        if line == "BEGIN:VEVENT":
            inside = True
            current_date = None
            current_summary = None
            continue
        if line == "END:VEVENT":
            if inside and current_date:
                events[current_date] = current_summary or "Feiertag"
            inside = False
            continue
        if not inside:
            continue
        if line.startswith("SUMMARY"):
            try:
                _, value = line.split(":", 1)
            except ValueError:
                continue
            current_summary = value.strip()
            continue
        if line.startswith("DTSTART"):
            try:
                _, value = line.split(":", 1)
            except ValueError:
                continue
            value = value.strip()
            parsed: Optional[dt.date] = None
            for fmt in ("%Y%m%d", "%Y%m%dT%H%M%SZ", "%Y%m%dT%H%M%S"):
                try:
                    parsed_dt = dt.datetime.strptime(value, fmt)
                except ValueError:
                    continue
                else:
                    parsed = parsed_dt.date()
                    break
            if parsed:
                current_date = parsed
    return sorted(events.items(), key=lambda item: item[0])


def import_holidays_from_ics(db: Session, content: str) -> List[Holiday]:
    parsed = _parse_ics_holidays(content)
    if not parsed:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="ICS-Datei enthält keine Feiertage")
    affected: Dict[dt.date, Holiday] = {}
    for day, name in parsed:
        holiday = db.query(Holiday).filter(Holiday.day == day).one_or_none()
        if holiday:
            holiday.name = name
            holiday.source = "ics"
            affected[day] = holiday
        else:
            holiday = Holiday(day=day, name=name, source="ics")
            db.add(holiday)
            affected[day] = holiday
    db.commit()
    for holiday in affected.values():
        db.refresh(holiday)
    return sorted(affected.values(), key=lambda item: item.day)


def create_calendar_event(
    db: Session,
    title: str,
    start_time: dt.datetime,
    end_time: dt.datetime,
    location: Optional[str],
    description: Optional[str],
    participated: bool,
    status_value: Optional[str],
    attendees: Optional[List[str]] = None,
) -> CalendarEvent:
    start_time_utc = _ensure_utc(start_time)
    end_time_utc = _ensure_utc(end_time)
    if end_time_utc <= start_time_utc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Event end must be after start")
    status_normalized = status_value or ("attended" if participated else "pending")
    if status_normalized not in CALENDAR_EVENT_STATUSES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unbekannter Kalenderstatus")
    ignored = status_normalized == "cancelled"
    effective_participated = status_normalized == "attended"
    event = CalendarEvent(
        title=title,
        start_time=start_time_utc,
        end_time=end_time_utc,
        location=location,
        description=description,
        participated=effective_participated,
        status=status_normalized,
        ignored=ignored,
        attendees=list(attendees or []),
        calendar_identifier="manual",
    )
    db.add(event)
    db.commit()
    db.refresh(event)
    return event


def list_calendar_events(
    db: Session,
    state: RuntimeState,
    start_date: Optional[dt.date],
    end_date: Optional[dt.date],
) -> List[CalendarEvent]:
    sync_caldav_events(db, state, start_date, end_date)
    query = db.query(CalendarEvent).filter(CalendarEvent.ignored.is_(False))
    if start_date:
        start, _ = _day_bounds(start_date)
        query = query.filter(CalendarEvent.start_time >= start)
    if end_date:
        _, end = _day_bounds(end_date)
        query = query.filter(CalendarEvent.start_time <= end)
    return query.order_by(CalendarEvent.start_time.asc()).all()


def set_calendar_participation(
    db: Session,
    state: RuntimeState,
    event_id: int,
    participated: Optional[bool] = None,
    status_value: Optional[str] = None,
    ignored: Optional[bool] = None,
) -> CalendarEvent:
    event = db.query(CalendarEvent).filter(CalendarEvent.id == event_id).one_or_none()
    if not event:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Calendar event not found")
    target_status = status_value
    if participated is not None:
        target_status = "attended" if participated else (status_value or "absent")
    if target_status is None:
        target_status = event.status or "pending"
    if target_status not in CALENDAR_EVENT_STATUSES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unbekannter Kalenderstatus")

    target_participated = target_status == "attended"
    target_ignored = ignored if ignored is not None else target_status == "cancelled"

    if (
        event.status == target_status
        and event.participated == target_participated
        and event.ignored == target_ignored
    ):
        if target_participated and event.subtrack is not None:
            _ensure_session_for_subtrack(db, state, event.subtrack, event.title)
        return event

    event.status = target_status
    event.participated = target_participated
    event.ignored = target_ignored

    if target_participated:
        start_local = _ensure_utc(event.start_time).astimezone(LOCAL_TZ)
        subtrack = event.subtrack
        if subtrack is None:
            subtrack = WorkSubtrack(
                day=start_local.date(),
                title=event.title,
                start_time=event.start_time,
                end_time=event.end_time,
                project=None,
                tags=[],
                note=event.description,
                calendar_event=event,
            )
            db.add(subtrack)
        else:
            subtrack.day = start_local.date()
            subtrack.title = event.title
            subtrack.start_time = event.start_time
            subtrack.end_time = event.end_time
            subtrack.note = event.description
        _ensure_session_for_subtrack(db, state, subtrack, event.title)
    else:
        subtrack = event.subtrack
        if subtrack is not None:
            day = subtrack.day
            auto_comment = _auto_session_comment(event.title, subtrack.note)
            session = (
                db.query(WorkSession)
                .filter(
                    WorkSession.start_time == event.start_time,
                    WorkSession.stop_time == event.end_time,
                    WorkSession.comment == auto_comment,
                )
                .one_or_none()
            )
            db.delete(subtrack)
            if session is not None:
                db.delete(session)
                if day:
                    update_day_summary(db, day, state)

    db.add(event)
    db.commit()
    db.refresh(event)
    return event


def update_runtime_settings(db: Session, state: RuntimeState, updates: dict) -> dict:
    normalized = dict(updates)
    if "block_ips" in normalized and normalized["block_ips"] is not None:
        if isinstance(normalized["block_ips"], str):
            normalized["block_ips"] = [ip.strip() for ip in normalized["block_ips"].split(",") if ip.strip()]
    if "caldav_selected_calendars" in normalized and normalized["caldav_selected_calendars"] is not None:
        if isinstance(normalized["caldav_selected_calendars"], str):
            normalized["caldav_selected_calendars"] = [
                entry.strip()
                for entry in normalized["caldav_selected_calendars"].split(",")
                if entry.strip()
            ]
        normalized["caldav_selected_calendars"] = normalize_calendar_selection(
            normalized["caldav_selected_calendars"]
        )
    if "caldav_default_cal" in normalized and normalized["caldav_default_cal"] is not None:
        normalized["caldav_default_cal"] = normalize_calendar_identifier(
            normalized["caldav_default_cal"]
        )
    state.apply(normalized)
    state.persist(
        db,
        {
            k: normalized[k]
            for k in normalized
            if k
            in {
                "block_ips",
                "caldav_url",
                "caldav_user",
                "caldav_password",
                "caldav_default_cal",
                "caldav_selected_calendars",
                "expected_daily_hours",
                "expected_weekly_hours",
                "vacation_days_per_year",
                "vacation_days_carryover",
                "day_overview_refresh_seconds",
                "travel_sender_contact",
                "travel_hr_contact",
                "travel_letter_template",
            }
        },
    )
    return state.snapshot()


def _write_pdf(path: Path, title: str, sessions: Iterable[WorkSession]) -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 2 * cm
    pdf.setTitle(title)
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(2 * cm, y, title)
    y -= 1 * cm
    pdf.setFont("Helvetica", 11)
    for session in sessions:
        line = f"{session.start_time.isoformat()} - {session.stop_time.isoformat() if session.stop_time else 'laufend'} | {(session.total_seconds or 0)/3600:.2f}h"
        pdf.drawString(2 * cm, y, line)
        y -= 0.8 * cm
        if y < 2 * cm:
            pdf.showPage()
            y = height - 2 * cm
            pdf.setFont("Helvetica", 11)
    pdf.save()


def _write_xlsx(path: Path, sessions: Iterable[WorkSession]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    ws.append(["Start", "Stop", "Dauer (h)", "Projekt", "Notiz"])
    for session in sessions:
        duration_hours = (session.total_seconds or 0) / 3600
        ws.append(
            [
                session.start_time.isoformat(),
                session.stop_time.isoformat() if session.stop_time else "",
                round(duration_hours, 2),
                session.project or "",
                session.comment or "",
            ]
        )
    wb.save(path)


def _seconds_to_hours(seconds: int) -> float:
    return round(seconds / 3600, 2)


def _collect_full_export_data(
    db: Session, state: RuntimeState, start_date: dt.date, end_date: dt.date
) -> List[Tuple[str, float, float]]:
    summaries = range_day_summaries(db, start_date, end_date, state)
    work_seconds = sum(item["work_seconds"] for item in summaries)
    vacation_seconds = sum(item["vacation_seconds"] for item in summaries)
    sick_seconds = sum(item["sick_seconds"] for item in summaries)
    work_days = sum(1 for item in summaries if item["work_seconds"] > 0)
    vacation_days = 0.0
    sick_days = 0.0
    for item in summaries:
        baseline = item.get("baseline_expected_seconds") or _expected_daily_seconds(state)
        if item["vacation_seconds"] > 0 and baseline:
            vacation_days += item["vacation_seconds"] / baseline
        if item["sick_seconds"] > 0 and baseline:
            sick_days += item["sick_seconds"] / baseline
    return [
        ("Arbeitstage", float(work_days), _seconds_to_hours(work_seconds)),
        ("Urlaubstage", round(vacation_days, 2), _seconds_to_hours(vacation_seconds)),
        ("AU-Tage", round(sick_days, 2), _seconds_to_hours(sick_seconds)),
    ]


def _write_full_pdf(path: Path, title: str, rows: List[Tuple[str, float, float]]) -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 2 * cm
    pdf.setTitle(title)
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(2 * cm, y, title)
    y -= 1.2 * cm
    pdf.setFont("Helvetica", 12)
    headers = ["Kategorie", "Tage", "Stunden"]
    column_widths = [8 * cm, 4 * cm, 4 * cm]
    pdf.drawString(2 * cm, y, headers[0])
    pdf.drawString(2 * cm + column_widths[0], y, headers[1])
    pdf.drawString(2 * cm + column_widths[0] + column_widths[1], y, headers[2])
    y -= 0.8 * cm
    pdf.setFont("Helvetica", 11)
    for category, days, hours in rows:
        pdf.drawString(2 * cm, y, category)
        pdf.drawRightString(2 * cm + column_widths[0] + column_widths[1] - 0.2 * cm, y, f"{days:.2f}")
        pdf.drawRightString(width - 2 * cm, y, f"{hours:.2f}")
        y -= 0.7 * cm
        if y < 2 * cm:
            pdf.showPage()
            y = height - 2 * cm
            pdf.setFont("Helvetica", 11)
    pdf.save()


def _write_full_xlsx(path: Path, rows: List[Tuple[str, float, float]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Zusammenfassung"
    ws.append(["Kategorie", "Tage", "Stunden"])
    for category, days, hours in rows:
        ws.append([category, float(f"{days:.2f}"), float(f"{hours:.2f}")])
    wb.save(path)


def export_sessions(
    db: Session,
    state: RuntimeState,
    export_type: str,
    export_format: str,
    start_date: dt.date,
    end_date: dt.date,
) -> ExportRecord:
    if export_type not in {"timesheet", "leave", "full"}:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported export type")
    if export_format not in {"pdf", "xlsx"}:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported export format")

    summary_rows: List[Tuple[str, float, float]] = []
    if export_type == "timesheet":
        start, _ = _day_bounds(start_date)
        _, end = _day_bounds(end_date)
        sessions = (
            db.query(WorkSession)
            .filter(
                and_(
                    WorkSession.start_time >= start,
                    WorkSession.start_time <= end,
                    WorkSession.status == "stopped",
                )
            )
            .order_by(WorkSession.start_time.asc())
            .all()
        )
    elif export_type == "leave":
        leaves = list_leaves(db, start_date, end_date, None)
        sessions = []
        for leave in leaves:
            session = WorkSession(
                start_time=dt.datetime.combine(leave["start_date"], dt.time.min),
                stop_time=dt.datetime.combine(leave["end_date"], dt.time.max),
                total_seconds=((leave["end_date"] - leave["start_date"]).days + 1) * 8 * 3600,
                project=leave["type"],
                comment=leave.get("comment"),
                status="stopped",
                tags=["leave"],
            )
            sessions.append(session)
    else:
        summary_rows = _collect_full_export_data(db, state, start_date, end_date)
        sessions = []

    file_suffix = "pdf" if export_format == "pdf" else "xlsx"
    filename = f"export_{export_type}_{start_date}_{end_date}_{int(_now().timestamp())}.{file_suffix}"
    path = settings.export_dir / filename

    if export_format == "pdf":
        if export_type == "full":
            _write_full_pdf(path, f"TimeTrack Vollexport {start_date} – {end_date}", summary_rows)
        else:
            _write_pdf(path, f"TimeTrack Export - {export_type}", sessions)
    else:
        if export_type == "full":
            _write_full_xlsx(path, summary_rows)
        else:
            _write_xlsx(path, sessions)

    checksum = _checksum_file(path)
    export = ExportRecord(
        type=export_type,
        format=export_format,
        range_start=start_date,
        range_end=end_date,
        path=str(path),
        checksum=checksum,
    )
    db.add(export)
    db.commit()
    db.refresh(export)
    return export


def _checksum_file(path: Path) -> str:
    import hashlib

    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _travel_directory(trip_id: int) -> Path:
    path = settings.travel_dir / str(trip_id)
    path.mkdir(parents=True, exist_ok=True)
    return path


def _sanitize_filename(name: str) -> str:
    normalized = re.sub(r"[^A-Za-z0-9_.-]", "_", name.strip()) if name else ""
    return normalized or "dokument"


def _format_contact_block(contact: Dict[str, str]) -> List[str]:
    lines: List[str] = []
    for key in ("name", "company", "department"):
        value = contact.get(key, "").strip()
        if value:
            lines.append(value)
    street = contact.get("street", "").strip()
    if street:
        lines.append(street)
    postal = contact.get("postal_code", "").strip()
    city = contact.get("city", "").strip()
    city_line = " ".join(part for part in (postal, city) if part)
    if city_line:
        lines.append(city_line)
    phone = contact.get("phone", "").strip()
    if phone:
        lines.append(f"Telefon: {phone}")
    email = contact.get("email", "").strip()
    if email:
        lines.append(email)
    return lines


def _build_letter_context(state: RuntimeState, trip: TravelTrip) -> Dict[str, str]:
    sender = dict(state.travel_sender_contact)
    hr = dict(state.travel_hr_contact)
    now_local = _now().astimezone(LOCAL_TZ)
    context: Dict[str, str] = defaultdict(str)  # type: ignore[assignment]

    for key, value in sender.items():
        context[f"sender_{key}"] = value
    for key, value in hr.items():
        context[f"hr_{key}"] = value

    context.update(
        sender_name=sender.get("name", ""),
        sender_company=sender.get("company", ""),
        sender_department=sender.get("department", ""),
        sender_street=sender.get("street", ""),
        sender_postal_code=sender.get("postal_code", ""),
        sender_city=sender.get("city", ""),
        sender_phone=sender.get("phone", ""),
        sender_email=sender.get("email", ""),
        hr_name=hr.get("name", ""),
        hr_company=hr.get("company", ""),
        hr_department=hr.get("department", ""),
        hr_street=hr.get("street", ""),
        hr_postal_code=hr.get("postal_code", ""),
        hr_city=hr.get("city", ""),
        hr_phone=hr.get("phone", ""),
        hr_email=hr.get("email", ""),
        today_date=now_local.strftime("%d.%m.%Y"),
        trip_title=trip.title,
        trip_destination=trip.destination or "",
        trip_purpose=trip.purpose or "",
        trip_start_date=trip.start_date.strftime("%d.%m.%Y"),
        trip_end_date=trip.end_date.strftime("%d.%m.%Y"),
    )

    duration_days = (trip.end_date - trip.start_date).days + 1
    context["trip_duration_days"] = str(duration_days)
    period = f"{trip.start_date.strftime('%d.%m.%Y')} – {trip.end_date.strftime('%d.%m.%Y')}"
    context["trip_period"] = period

    destination = trip.destination or ""
    if destination:
        context["trip_destination_clause"] = f"nach {destination}"
    else:
        context["trip_destination_clause"] = ""

    purpose = trip.purpose or ""
    if purpose:
        context["trip_purpose_clause"] = f"wegen {purpose}"
    else:
        context["trip_purpose_clause"] = ""

    sender_block = "\n".join(_format_contact_block(sender)).strip()
    hr_block = "\n".join(_format_contact_block(hr)).strip()
    context["sender_block"] = sender_block
    context["hr_block"] = hr_block

    return context


def _render_letter_template(template: str, context: Dict[str, str]) -> str:
    if not template:
        return ""
    safe_context: defaultdict[str, str] = defaultdict(str)  # type: ignore[assignment]
    safe_context.update(context)
    try:
        rendered = template.format_map(safe_context)
    except Exception:
        return template
    return rendered


def generate_travel_letter_preview(db: Session, state: RuntimeState, trip_id: int) -> Dict[str, Any]:
    trip = _get_travel_trip(db, trip_id)
    context = _build_letter_context(state, trip)
    template = state.travel_letter_template
    subject = _render_letter_template(template.get("subject", ""), context).strip()
    body = _render_letter_template(template.get("body", ""), context)
    body = body.replace("\r\n", "\n").strip()
    return {
        "subject": subject or template.get("subject", ""),
        "body": body or template.get("body", ""),
        "context": dict(context),
        "sender_contact": dict(state.travel_sender_contact),
        "hr_contact": dict(state.travel_hr_contact),
    }


def _render_travel_letter_pdf(
    trip: TravelTrip,
    subject: str,
    body: str,
    sender_contact: Dict[str, str],
    hr_contact: Dict[str, str],
) -> bytes:
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    page.setTitle(f"Anschreiben {trip.title}")
    top_margin = height - 2.5 * cm
    y = top_margin

    def ensure_space(current_y: float, required: float, font: tuple[str, int]) -> float:
        if current_y - required < 2 * cm:
            page.showPage()
            page.setFont(font[0], font[1])
            return height - 2 * cm
        return current_y

    page.setFont("Helvetica", 9)
    for line in _format_contact_block(sender_contact):
        if not line:
            continue
        y = ensure_space(y, 0.5 * cm, ("Helvetica", 9))
        page.drawString(2 * cm, y, line)
        y -= 0.5 * cm

    y -= 0.7 * cm
    page.setFont("Helvetica", 11)
    for line in _format_contact_block(hr_contact):
        if not line:
            continue
        y = ensure_space(y, 0.6 * cm, ("Helvetica", 11))
        page.drawString(2 * cm, y, line)
        y -= 0.6 * cm

    now_local = _now().astimezone(LOCAL_TZ)
    date_label = now_local.strftime("%d.%m.%Y")
    y -= 0.6 * cm
    page.setFont("Helvetica", 11)
    page.drawString(2 * cm, y, date_label)
    y -= 1.0 * cm

    subject_text = subject.strip()
    if subject_text:
        page.setFont("Helvetica-Bold", 12)
        page.drawString(2 * cm, y, f"Betreff: {subject_text}")
        y -= 1.2 * cm

    page.setFont("Helvetica", 11)
    line_height = 0.65 * cm
    paragraphs = body.replace("\r\n", "\n").split("\n\n") if body else []

    for paragraph in paragraphs:
        lines = paragraph.split("\n") if paragraph else [""]
        for line in lines:
            chunks = _wrap_text(line.strip(), 88) or [""]
            for chunk in chunks:
                y = ensure_space(y, line_height, ("Helvetica", 11))
                page.drawString(2 * cm, y, chunk)
                y -= line_height
        y -= 0.4 * cm

    page.save()
    buffer.seek(0)
    return buffer.read()


def create_travel_letter_document(
    db: Session,
    state: RuntimeState,
    trip_id: int,
    subject: str,
    body: str,
) -> TravelDocument:
    cleaned_subject = subject.strip()
    cleaned_body = body.replace("\r\n", "\n").strip()
    if not cleaned_subject or not cleaned_body:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Betreff und Text dürfen nicht leer sein")
    trip = _get_travel_trip(db, trip_id)
    pdf_bytes = _render_travel_letter_pdf(
        trip,
        cleaned_subject,
        cleaned_body,
        state.travel_sender_contact,
        state.travel_hr_contact,
    )
    timestamp = _now().astimezone(LOCAL_TZ).strftime("%Y%m%d_%H%M%S")
    original_name = f"Anschreiben_{trip.start_date}_{trip.id}_{timestamp}.pdf"
    comment = f"Automatisch erstellt am {dt.datetime.now(LOCAL_TZ).strftime('%d.%m.%Y %H:%M')}"
    document = add_travel_document(
        db,
        trip.id,
        "Anschreiben",
        original_name,
        pdf_bytes,
        comment,
    )
    return document


def _get_travel_trip(db: Session, trip_id: int) -> TravelTrip:
    trip = db.query(TravelTrip).filter(TravelTrip.id == trip_id).one_or_none()
    if not trip:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Dienstreise nicht gefunden")
    return trip


def list_travel_trips(db: Session) -> List[TravelTrip]:
    return (
        db.query(TravelTrip)
        .order_by(TravelTrip.start_date.desc(), TravelTrip.id.desc())
        .all()
    )


def create_travel_trip(
    db: Session,
    title: str,
    start_date: dt.date,
    end_date: dt.date,
    destination: Optional[str],
    purpose: Optional[str],
    workflow_state: Optional[str],
    notes: Optional[str],
) -> TravelTrip:
    if end_date < start_date:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Enddatum liegt vor dem Startdatum")
    state_value = workflow_state or "request_draft"
    if state_value not in TRAVEL_WORKFLOW_STATES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unbekannter Workflow-Status")
    trip = TravelTrip(
        title=title,
        start_date=start_date,
        end_date=end_date,
        destination=destination,
        purpose=purpose,
        workflow_state=state_value,
        notes=notes,
    )
    db.add(trip)
    db.commit()
    db.refresh(trip)
    _travel_directory(trip.id)
    return trip


def update_travel_trip(
    db: Session,
    trip_id: int,
    *,
    title: Optional[str] = None,
    start_date: Optional[dt.date] = None,
    end_date: Optional[dt.date] = None,
    destination: Optional[str] = None,
    purpose: Optional[str] = None,
    workflow_state: Optional[str] = None,
    notes: Optional[str] = None,
) -> TravelTrip:
    trip = _get_travel_trip(db, trip_id)
    new_start = start_date or trip.start_date
    new_end = end_date or trip.end_date
    if new_end < new_start:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Enddatum liegt vor dem Startdatum")
    if workflow_state is not None and workflow_state not in TRAVEL_WORKFLOW_STATES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unbekannter Workflow-Status")

    if title is not None:
        trip.title = title
    trip.start_date = new_start
    trip.end_date = new_end
    if destination is not None:
        trip.destination = destination
    if purpose is not None:
        trip.purpose = purpose
    if workflow_state is not None:
        trip.workflow_state = workflow_state
    if notes is not None:
        trip.notes = notes

    db.add(trip)
    db.commit()
    db.refresh(trip)
    return trip


def delete_travel_trip(db: Session, trip_id: int) -> None:
    trip = _get_travel_trip(db, trip_id)
    directory = settings.travel_dir / str(trip.id)
    db.delete(trip)
    db.commit()
    if directory.exists():
        shutil.rmtree(directory, ignore_errors=True)


def add_travel_document(
    db: Session,
    trip_id: int,
    document_type: str,
    original_name: str,
    content: bytes,
    comment: Optional[str],
    collection_label: Optional[str] = None,
    linked_invoice_id: Optional[int] = None,
) -> TravelDocument:
    if document_type not in TRAVEL_DOCUMENT_TYPES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unbekannter Dokumenttyp")
    trip = _get_travel_trip(db, trip_id)
    directory = _travel_directory(trip.id)
    suffix = Path(original_name).suffix if original_name else ""
    stored_name = f"{uuid.uuid4().hex}{suffix}"
    stored_path = directory / stored_name
    with stored_path.open("wb") as handle:
        handle.write(content)
    normalized_label = _normalize_collection_label(collection_label)
    linked_invoice: Optional[TravelDocument] = None
    if linked_invoice_id is not None:
        invoice = _get_travel_document(db, trip_id, linked_invoice_id)
        if invoice.document_type != "Rechnung":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Verknüpfungen sind nur mit Rechnungen möglich",
            )
        if document_type == "Rechnung":
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Rechnungen können nicht mit anderen Rechnungen verknüpft werden",
            )
        linked_invoice = invoice
    next_index = (
        db.query(func.max(TravelDocument.sort_index))
        .filter(TravelDocument.trip_id == trip.id)
        .scalar()
    )
    sort_index = (next_index or 0) + 1 if next_index is not None else 0
    document = TravelDocument(
        trip=trip,
        document_type=document_type,
        stored_path=str(stored_path),
        original_name=original_name or stored_name,
        comment=comment,
        collection_label=normalized_label,
        sort_index=sort_index,
    )
    if linked_invoice is not None:
        document.linked_invoice = linked_invoice
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


def _get_travel_document(db: Session, trip_id: int, document_id: int) -> TravelDocument:
    document = (
        db.query(TravelDocument)
        .filter(TravelDocument.id == document_id, TravelDocument.trip_id == trip_id)
        .one_or_none()
    )
    if not document:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Dokument nicht gefunden")
    return document


def update_travel_document(
    db: Session,
    trip_id: int,
    document_id: int,
    *,
    comment: Any = UNSET,
    signed: Any = UNSET,
    collection_label: Any = UNSET,
    linked_invoice_id: Any = UNSET,
) -> TravelDocument:
    document = _get_travel_document(db, trip_id, document_id)
    if comment is not UNSET:
        document.comment = comment
    if signed is not UNSET:
        if document.document_type not in SIGNABLE_DOCUMENT_TYPES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Dokument kann nicht signiert werden")
        document.signed = signed
    if collection_label is not UNSET:
        document.collection_label = _normalize_collection_label(collection_label)
    if linked_invoice_id is not UNSET:
        if document.document_type == "Rechnung" and linked_invoice_id not in (None, document.id):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Rechnungen können nicht mit anderen Rechnungen verknüpft werden",
            )
        if linked_invoice_id is None:
            document.linked_invoice = None
        else:
            invoice = _get_travel_document(db, trip_id, linked_invoice_id)
            if invoice.document_type != "Rechnung":
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Verknüpfungen sind nur mit Rechnungen möglich",
                )
            if invoice.id == document.id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Ein Dokument kann nicht mit sich selbst verknüpft werden",
                )
            document.linked_invoice = invoice
    db.add(document)
    db.commit()
    db.refresh(document)
    return document


def reorder_travel_documents(db: Session, trip_id: int, ordered_ids: List[int]) -> TravelTrip:
    trip = _get_travel_trip(db, trip_id)
    existing_ids = [document.id for document in trip.documents]
    if not ordered_ids:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Es wurden keine Dokumente zur Neuanordnung übermittelt",
        )
    if len(ordered_ids) != len(existing_ids) or set(ordered_ids) != set(existing_ids):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Die übermittelte Reihenfolge stimmt nicht mit den vorhandenen Dokumenten überein",
        )
    documents_by_id = {document.id: document for document in trip.documents}
    for position, document_id in enumerate(ordered_ids):
        document = documents_by_id.get(document_id)
        if document is None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Unbekanntes Dokument in der Reihenfolge entdeckt",
            )
        document.sort_index = position
        db.add(document)
    db.commit()
    db.refresh(trip)
    db.expire(trip, ["documents"])
    _ = trip.documents  # force reload with updated order
    return trip


def delete_travel_document(db: Session, trip_id: int, document_id: int) -> None:
    document = _get_travel_document(db, trip_id, document_id)
    stored_path = Path(document.stored_path)
    if document.document_type == "Rechnung":
        linked = (
            db.query(TravelDocument)
            .filter(TravelDocument.trip_id == trip_id, TravelDocument.linked_invoice_id == document.id)
            .all()
        )
        for receipt in linked:
            receipt.linked_invoice = None
            db.add(receipt)
    db.delete(document)
    db.commit()
    if stored_path.exists():
        try:
            stored_path.unlink()
        except OSError:
            pass


def resolve_travel_document_path(db: Session, trip_id: int, document_id: int) -> Tuple[TravelDocument, Path]:
    document = _get_travel_document(db, trip_id, document_id)
    stored_path = Path(document.stored_path)
    if not stored_path.exists():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Datei wurde nicht gefunden")
    return document, stored_path


IMAGE_EXTENSIONS: Set[str] = {".png", ".jpg", ".jpeg"}


def _wrap_text(text: str, width: int) -> List[str]:
    words = text.split()
    if not words:
        return [""]
    lines: List[str] = []
    current = words[0]
    for word in words[1:]:
        tentative = f"{current} {word}"
        if len(tentative) > width:
            lines.append(current)
            current = word
        else:
            current = tentative
    lines.append(current)
    return lines


def _collect_travel_dataset_documents(db: Session, trip_id: int) -> Tuple[TravelTrip, List[TravelDocument]]:
    trip = _get_travel_trip(db, trip_id)
    documents = (
        db.query(TravelDocument)
        .filter(
            TravelDocument.trip_id == trip.id,
            TravelDocument.document_type.in_(TRAVEL_DATASET_DOCUMENT_TYPES),
        )
        .order_by(
            TravelDocument.sort_index.asc(),
            TravelDocument.created_at.asc(),
            TravelDocument.id.asc(),
        )
        .all()
    )
    return trip, documents


def _render_dataset_cover_page(trip: TravelTrip, documents: List[TravelDocument]) -> io.BytesIO:
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    page.setTitle("Reisekostenabrechnungsdatensatz")
    page.setFont("Helvetica-Bold", 18)
    page.drawString(2 * cm, height - 2.5 * cm, "Reisekostenabrechnungsdatensatz")
    page.setFont("Helvetica", 11)
    lines = [
        f"Dienstreise: {trip.title}",
        f"Zeitraum: {trip.start_date.strftime('%d.%m.%Y')} – {trip.end_date.strftime('%d.%m.%Y')}",
    ]
    if trip.destination:
        lines.append(f"Zielort: {trip.destination}")
    if trip.purpose:
        lines.append(f"Reisezweck: {trip.purpose}")
    y = height - 4 * cm
    for entry in lines:
        page.drawString(2 * cm, y, entry)
        y -= 0.8 * cm
    page.setFont("Helvetica-Bold", 12)
    page.drawString(2 * cm, y - 0.5 * cm, "Enthaltene Dokumente")
    y -= 1.5 * cm
    text = page.beginText(2 * cm, y)
    text.setFont("Helvetica", 10)
    if documents:
        for document in documents:
            parts = [f"{document.document_type}: {document.original_name}"]
            if document.collection_label:
                parts.insert(0, f"[{document.collection_label}]")
            if document.linked_invoice and document.document_type != "Rechnung":
                parts.append(f"(Rechnung: {document.linked_invoice.original_name})")
            label = "- " + " ".join(parts)
            for wrapped in _wrap_text(label, 88):
                text.textLine(wrapped)
    else:
        text.textLine("Keine Dokumente vorhanden")
    page.drawText(text)
    page.showPage()
    page.save()
    buffer.seek(0)
    return buffer


def _render_dataset_image_page(document: TravelDocument, path: Path) -> io.BytesIO:
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    page.setFont("Helvetica-Bold", 13)
    page.drawString(2 * cm, height - 2.5 * cm, f"{document.document_type}: {document.original_name}")
    page.setFont("Helvetica", 9)
    info_y = height - 3.2 * cm
    page.drawString(2 * cm, info_y, f"Datei: {path.name}")
    info_y -= 0.5 * cm
    if document.collection_label:
        page.drawString(2 * cm, info_y, f"Sammelbegriff: {document.collection_label}")
        info_y -= 0.5 * cm
    if document.linked_invoice and document.document_type != "Rechnung":
        page.drawString(2 * cm, info_y, f"Verknüpfte Rechnung: {document.linked_invoice.original_name}")
        info_y -= 0.5 * cm
    image = ImageReader(str(path))
    img_width, img_height = image.getSize()
    if img_width <= 0 or img_height <= 0:
        raise ValueError("Ungültige Bildabmessungen")
    available_width = width - 4 * cm
    top_limit = max(info_y - 0.5 * cm, 6 * cm)
    bottom_margin = 2 * cm
    available_height = max(top_limit - bottom_margin, 4 * cm)
    scale = min(available_width / img_width, available_height / img_height, 1.0)
    draw_width = img_width * scale
    draw_height = img_height * scale
    x = (width - draw_width) / 2
    y = bottom_margin + max((available_height - draw_height) / 2, 0)
    page.drawImage(image, x, y, width=draw_width, height=draw_height, preserveAspectRatio=True, mask="auto")
    page.showPage()
    page.save()
    buffer.seek(0)
    return buffer


def _render_dataset_placeholder_page(document: TravelDocument, message: str) -> io.BytesIO:
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    page.setFont("Helvetica-Bold", 14)
    page.drawString(2 * cm, height - 2.5 * cm, document.original_name)
    page.setFont("Helvetica", 10)
    text = page.beginText(2 * cm, height - 4 * cm)
    text.textLine(f"Dokumenttyp: {document.document_type}")
    text.textLine(f"Originaldatei: {document.original_name}")
    if document.collection_label:
        text.textLine(f"Sammelbegriff: {document.collection_label}")
    if document.linked_invoice and document.document_type != "Rechnung":
        text.textLine(f"Verknüpfte Rechnung: {document.linked_invoice.original_name}")
    text.textLine("")
    for wrapped in _wrap_text(message, 90):
        text.textLine(wrapped)
    page.drawText(text)
    page.showPage()
    page.save()
    buffer.seek(0)
    return buffer


def build_travel_dataset_archive(db: Session, trip_id: int) -> Tuple[str, bytes]:
    trip, documents = _collect_travel_dataset_documents(db, trip_id)
    if not documents:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Keine abrechnungsrelevanten Dokumente vorhanden",
        )
    buffer = io.BytesIO()
    added = 0
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for document in documents:
            stored_path = Path(document.stored_path)
            if not stored_path.exists():
                continue
            safe_name = _sanitize_filename(document.original_name)
            prefix = _sanitize_filename(document.document_type)
            arcname = f"{prefix}_{safe_name}" if prefix else safe_name
            archive.write(stored_path, arcname)
            added += 1
    if added == 0:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Dokumentdateien fehlen")
    buffer.seek(0)
    filename = f"reisekosten_{trip.start_date}_{trip.end_date}.zip"
    return filename, buffer.read()


def build_travel_dataset_pdf(db: Session, trip_id: int) -> Tuple[str, bytes]:
    trip, documents = _collect_travel_dataset_documents(db, trip_id)
    if not documents:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Keine abrechnungsrelevanten Dokumente vorhanden",
        )

    merger = PdfMerger()
    pages_added = 0

    cover_page = _render_dataset_cover_page(trip, documents)
    merger.append(cover_page)
    pages_added += 1

    for document in documents:
        stored_path = Path(document.stored_path)
        if not stored_path.exists():
            placeholder = _render_dataset_placeholder_page(
                document,
                "Die zugehörige Originaldatei konnte nicht gefunden werden und ist daher nicht im Ausdruck enthalten.",
            )
            merger.append(placeholder)
            pages_added += 1
            continue

        suffix = stored_path.suffix.lower()
        if suffix == ".pdf":
            merger.append(str(stored_path))
            pages_added += 1
        elif suffix in IMAGE_EXTENSIONS:
            try:
                image_page = _render_dataset_image_page(document, stored_path)
            except Exception:  # pragma: no cover - fallback for broken images
                placeholder = _render_dataset_placeholder_page(
                    document,
                    "Die Bilddatei konnte nicht verarbeitet werden. Bitte laden Sie die Originaldatei separat herunter.",
                )
                merger.append(placeholder)
            else:
                merger.append(image_page)
            pages_added += 1
        else:
            placeholder = _render_dataset_placeholder_page(
                document,
                "Dieses Dateiformat kann nicht automatisch eingebettet werden. Bitte laden Sie die Datei separat herunter.",
            )
            merger.append(placeholder)
            pages_added += 1

    if pages_added == 0:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Keine druckbaren Inhalte gefunden")

    output = io.BytesIO()
    merger.write(output)
    merger.close()
    output.seek(0)
    filename = f"reisekosten_{trip.start_date}_{trip.end_date}.pdf"
    return filename, output.read()
